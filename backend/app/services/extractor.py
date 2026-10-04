import io
from fastapi import UploadFile, HTTPException
from pypdf import PdfReader
from docx import Document

async def extract_text_from_upload(file: UploadFile) -> str:
    """Extract clean text content from PDF, DOCX, or plain text files."""
    filename = file.filename or ""
    content = await file.read()
    
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    
    ext = filename.lower().split(".")[-1] if "." in filename else ""
    
    try:
        if ext == "pdf":
            reader = PdfReader(io.BytesIO(content))
            pages_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            extracted = "\n".join(pages_text).strip()
            if not extracted:
                raise ValueError("Could not extract readable text from PDF (it may be a scanned image).")
            return extracted
            
        elif ext in ["docx", "doc"]:
            doc = Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            return "\n".join(paragraphs).strip()
            
        elif ext in ["txt", "md", "rtf", "json"]:
            try:
                return content.decode("utf-8").strip()
            except UnicodeDecodeError:
                return content.decode("latin-1", errors="ignore").strip()
        else:
            # Fallback attempt text decode
            try:
                return content.decode("utf-8").strip()
            except Exception:
                raise HTTPException(
                    status_code=400,
                    detail=f"Unsupported file format '{ext}'. Please upload PDF, DOCX, or TXT file."
                )
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Failed to process file {filename}: {str(e)}"
        )
